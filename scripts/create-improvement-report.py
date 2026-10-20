#!/usr/bin/env python3
"""Build a separate improvement report from an immutable offline snapshot.

No provider calls, benchmark mutation, or application writes. Use the bundled
document Python runtime; render and inspect all pages before delivery.
"""
import argparse
import collections
import datetime as dt
import hashlib
import json
import pathlib

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def pct(value):
    return "Not available" if value is None else f"{value * 100:.2f}%"


def put(parent, tag, attrs=None):
    element = OxmlElement(tag)
    for key, value in (attrs or {}).items():
        element.set(qn(key), str(value))
    parent.append(element)
    return element


def create(snapshot, output):
    manifest = json.loads((snapshot / "manifest.json").read_text())
    for item in manifest["files"]:
        assert hashlib.sha256((snapshot / item["name"]).read_bytes()).hexdigest() == item["sha256"], item["name"]
    report = json.loads((snapshot / "content-agreement-report.json").read_text())
    raw = json.loads((snapshot / "candidate-exact-report.json").read_text())
    design = json.loads((snapshot / "evaluation-design.json").read_text())
    reviews = json.loads((snapshot / "content-adjudications.json").read_text())
    suggestions = json.loads((snapshot / "goods-description-human-review-suggestions.json").read_text())
    b, c = report["baseline"], report["candidate"]
    bm, cm = b["developmentOverall"], c["developmentOverall"]
    bv, cv = b["visibleUiComparableSubset"], c["visibleUiComparableSubset"]
    dev_ids = {p["id"] for p in design["development"]}
    packets = [p for p in raw["packets"] if p["packetId"] in dev_ids]
    completed = sum(p["readingStatus"] == "complete" for p in packets)
    conflict_count = sum(len(p.get("labelConflicts", [])) for p in packets)
    extra = sum(p.get("extraFilledComponents", 0) for p in packets)
    extra_rows = sum(p.get("extraGoodsRows", 0) for p in packets)
    needed = max(0, -(-cv["eligible"] * 90 // 100) - cv["contentCorrect"])
    moment = dt.datetime.fromisoformat(report["generatedAt"].replace("Z", "+00:00")).strftime("%d %B %Y at %H:%M UTC")

    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.top_margin, s.bottom_margin = Inches(.65), Inches(.62)
    s.left_margin, s.right_margin = Inches(.75), Inches(.75)
    s.footer_distance = Inches(.28)
    for border in d.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = d.styles[name]
        style.font.name, style.font.color.rgb, style.font.underline = "Calibri", RGBColor(0, 0, 0), False
    d.styles["Normal"].font.size = Pt(10.5)
    d.styles["Normal"].paragraph_format.line_spacing = 1.06
    d.styles["Normal"].paragraph_format.space_after = Pt(7)
    for name, size in (("Title", 25), ("Heading 1", 19), ("Heading 2", 12)):
        style = d.styles[name]
        style.font.size, style.font.bold = Pt(size), True
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(7)
    d.core_properties.title = "AICD accuracy improvement report"
    d.core_properties.subject = "Fixed development comparison with explicit content review"
    d.core_properties.author = "AICD"
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("AICD improvement report  |  ").font.size = Pt(9)
    put(footer._p, "w:fldSimple", {"w:instr": "PAGE"})

    def p(text, lead=None, small=False):
        para = d.add_paragraph()
        if lead:
            para.add_run(lead + " ").bold = True
        para.add_run(text)
        if small:
            for run in para.runs:
                run.font.size = Pt(9)
        return para

    def heading(text, level=2):
        return d.add_heading(text, level)

    def page(text):
        para = heading(text, 1)
        para.paragraph_format.page_break_before = True

    def table(headers, rows, widths, size=10):
        t = d.add_table(rows=1, cols=len(headers))
        t.alignment, t.autofit = WD_TABLE_ALIGNMENT.CENTER, False
        for col, width in zip(t.columns, widths):
            col.width = Inches(width)
        borders = put(t._tbl.tblPr, "w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            put(borders, "w:" + edge, {"w:val": "single", "w:sz": 4, "w:color": "D9D9D9"})
        for i, row in enumerate([headers] + rows):
            cells = t.rows[0].cells if i == 0 else t.add_row().cells
            put(t.rows[i]._tr.get_or_add_trPr(), "w:cantSplit")
            if i == 0:
                put(t.rows[i]._tr.get_or_add_trPr(), "w:tblHeader")
            for j, (cell, text, width) in enumerate(zip(cells, row, widths)):
                cell.width, cell.vertical_alignment = Inches(width), WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr = cell._tc.get_or_add_tcPr()
                margins = put(pr, "w:tcMar")
                for edge, margin in (("top", 85), ("bottom", 85), ("left", 105), ("right", 105)):
                    put(margins, "w:" + edge, {"w:w": margin, "w:type": "dxa"})
                put(pr, "w:shd", {"w:fill": "DCE6EE" if i == 0 else "FFFFFF"})
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.02
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run(str(text)); run.font.size = Pt(size); run.bold = i == 0
        d.add_paragraph().paragraph_format.space_after = Pt(0)
        return t

    d.add_paragraph("AICD accuracy improvement report", "Title")
    p("Fixed development comparison against employee ECUS declarations")
    p("Results frozen on " + moment, small=True)
    p(f"The 90% target has not been met. Across comparable components in the three visible ECUS tabs, the improved draft achieves {cv['contentCorrect']:,}/{cv['eligible']:,} content agreement ({pct(cv['contentAgreement'])}) and {cv['exactMatches']:,}/{cv['eligible']:,} exact agreement ({pct(cv['exactAgreement'])}). This mapped subset does not cover every one of the 191 controls. All scored development components, including 232 additional structured goods components, reach {pct(cm['contentAgreement'])} content and {pct(cm['exactAgreement'])} exact agreement. These are development results, not independent validation or approved forms.")
    heading("Results on the same saved readings")
    table(["Metric", "Baseline", "Improved candidate"], [
        ["Visible tabs exact agreement", f"{bv['exactMatches']} / {bv['eligible']}\n{pct(bv['exactAgreement'])}", f"{cv['exactMatches']} / {cv['eligible']}\n{pct(cv['exactAgreement'])}"],
        ["Visible tabs content agreement", f"{bv['contentCorrect']} / {bv['eligible']}\n{pct(bv['contentAgreement'])}", f"{cv['contentCorrect']} / {cv['eligible']}\n{pct(cv['contentAgreement'])}"],
        ["Visible tabs filled content agreement", f"{bv['contentCorrect']} / {bv['filled']}\n{pct(bv['filledContentAgreement'])}", f"{cv['contentCorrect']} / {cv['filled']}\n{pct(cv['filledContentAgreement'])}"],
        ["Visible tabs comparable field coverage", f"{bv['filled']} / {bv['eligible']}\n{pct(bv['filledCoverage'])}", f"{cv['filled']} / {cv['eligible']}\n{pct(cv['filledCoverage'])}"],
        ["All scored development components exact agreement", f"{bm['exactMatches']} / {bm['eligible']}\n{pct(bm['exactAgreement'])}", f"{cm['exactMatches']} / {cm['eligible']}\n{pct(cm['exactAgreement'])}"],
        ["All scored development components content agreement", f"{bm['contentCorrect']} / {bm['eligible']}\n{pct(bm['contentAgreement'])}", f"{cm['contentCorrect']} / {cm['eligible']}\n{pct(cm['contentAgreement'])}"],
    ], [3.5, 1.75, 1.75])
    p(f"The comparison replays the same 200 saved page outcomes (198 successful provider responses and two saved provider failures) across the same 18 development collections and {cm['eligible']:,} answer components. There were no new Kimi calls for these product changes. {completed} collections have complete validated readings; {len(packets)-completed} retain provisional drafts. The 11 held-out eligible collections have not yet been evaluated for this improvement experiment.")
    p(f"The visible subset includes {cv['missing']:,} missing values and {cv['unreviewedMismatches'] + cv['reviewedNotEquivalent']:,} uncredited populated differences. Reaching 90% on this subset requires {needed:,} more credited components. All scored development components include {cm['missing']:,} missing values; its filled content agreement is {pct(cm['filledContentAgreement'])} and coverage is {pct(cm['filledCoverage'])}. Missing values count against agreement and coverage in both scopes.", lead="What remains.")

    page("Agreement across the three ECUS tabs")
    p("These figures map the existing answer components to the exact three tabs the employee fills. Rates compare individual components, not the percentage of wholly correct forms. The complete original denominator is retained.")
    rows = []
    for tab, label in (("general_1", "Thông tin chung"), ("general_2", "Thông tin chung 2"), ("goods", "Danh sách hàng"), ("structured_goods_only", "Structured goods\noutside visible controls")):
        x, y = b["byForm"][tab], c["byForm"][tab]
        rows.append([label, y["eligible"], pct(x["exactAgreement"]), pct(x["contentAgreement"]), pct(y["exactAgreement"]), pct(y["contentAgreement"])])
    table(["Area", "Eligible", "Baseline\nexact", "Baseline\ncontent", "Candidate\nexact", "Candidate\ncontent"], rows, [2.05, .6, 1.05, 1.1, 1.1, 1.1], 9)
    p("The 232 structured-only components are goods unit-price currency and unit values printed in ToKhai and stored in the draft. They remain in the overall denominator, but are shown separately because they are not separate controls in the visible form catalog.")
    heading("Complete readings still leave form gaps")
    bc, cc = b["completeReadingOnly"]["overall"], c["completeReadingOnly"]["overall"]
    table(["Four collections with complete readings", "Baseline", "Candidate"], [
        ["Exact agreement", f"{bc['exactMatches']} / {bc['eligible']}  ({pct(bc['exactAgreement'])})", f"{cc['exactMatches']} / {cc['eligible']}  ({pct(cc['exactAgreement'])})"],
        ["Content agreement", f"{bc['contentCorrect']} / {bc['eligible']}  ({pct(bc['contentAgreement'])})", f"{cc['contentCorrect']} / {cc['eligible']}  ({pct(cc['contentAgreement'])})"],
        ["Missing comparable components", bc["missing"], cc["missing"]],
    ], [3.4, 1.8, 1.8])
    p("This completed-reading table uses the full original 340-component denominator, including structured goods values outside visible controls. Completed readings mean the saved pages passed reading validation; they do not mean forms are complete, correct or approved. The headline keeps all 18 development collections, including partial drafts.")
    heading("The catalog is broader than the score")
    p("The three-tab catalog contains 191 controls, including repeatable structured entries. The current comparison does not validate every control, all taxes or generated customs identifiers. A separate coverage audit identifies populated print areas outside this score, including exporter address parts, remarks, attachment details, taxable values, duty amounts and VAT. These require defined inputs and validation before they can be added to a score.")

    page("Unresolved content and business context")
    p("Product fixes recover source-backed values that were lost during date, numeric, unit and draft-state handling. They do not authorize inventing customs decisions or copying the employee answer into the draft. Remaining missing values and disagreements need evidence from original sources or separately approved business reference data.")
    field_names = {"1.74":"HS classification", "1.78":"Goods description", "1.79":"Origin", "1.80":"Import duty schedule", "1.83":"Goods quantity and unit", "1.86":"Goods unit price", "1.20":"Exporter", "1.33":"Unloading location", "1.34":"Loading location", "1.43":"Payment method"}
    groups = collections.defaultdict(collections.Counter)
    for row in c["components"]:
        if row["contentStatus"] not in ("exact_match", "reviewed_equivalent"):
            groups[row["fieldId"]]["missing" if row["contentStatus"] == "missing" else "difference"] += 1
    top = sorted(groups.items(), key=lambda x: -sum(x[1].values()))[:7]
    table(["Largest unresolved component groups", "Missing", "Uncredited differences"], [[field_names.get(k, "ECUS field " + k), v["missing"], v["difference"]] for k,v in top], [4.0, 1.25, 1.75], 10)
    p(f"All {suggestions['descriptionMismatchCount']} description mismatches remain uncredited. Internal triage flags 86 for application or function-detail review, 44 for numeric references missing from the description text, five for condition or specification differences and seven for possible material identifier differences. Categories overlap; a missing description detail may exist elsewhere in the draft or original sources.")
    p("For example, an employee description includes diesel operation, power above 100 kW and used condition, while the draft identifies only the power-pack model. Other pairs differ in a filter model or serial digit. Goods are compared in printed row order; apparent item differences may arise from alignment, employee revisions or extraction. They are not automatically equivalent translations or proven AI errors.")
    heading("Documented equivalent representations")
    rules = collections.Counter(entry.get("ruleId", "Other explicit review") for entry in reviews["entries"] if entry["status"] == "equivalent")
    p(f"The review file contains {sum(rules.values())} explicit equivalent-pair decisions; {cm['reviewedEquivalent']} currently add candidate credit. Others apply to baseline or earlier pairs, including units that now match exactly. Credited differences cover company punctuation, the same city name with a generic qualifier, and historically documented literal payment methods. Only matching component identities and exact expected/actual pair hashes receive credit. Changed pairs require fresh review.")
    p("These decisions are AI-assisted internal review, not independent human validation. Missing values, unresolved translations, wrong numbers, HS codes, origins, specific-port differences, payment terms and unsupported unit substitutions receive no automatic semantic credit. A reviewed name or unit does not validate the rest of its field or shipment.")

    page("Evidence boundaries and next validation")
    heading("What was held fixed")
    p("Both versions use the same saved provider outcomes, original commercial source inputs, development IDs and employee answer components. The approved Kimi batch uses model kimi-k2.6 with 223 original files and 497 rendered pages and was still running when this report was prepared. ToKhai files remain scoring-only answer keys and are never sent to Kimi. The unchanged exact scorer remains visible beside supplementary content review. The audited source-folder partition and import/export direction are supplied fixture inputs; email intake and automatic shipment grouping are outside this test.")
    p(f"The complete sample inventory has 44 collections, of which 29 have supported eligible import answers. Four collections have no primary sources, nine are non-import and two lack a supported import answer key. This paired comparison uses only the 18 development collections. The other 11 eligible collections remain held out. Within development, {conflict_count} conflicting employee-revision components are excluded rather than resolved to whichever answer is easier to match. Other unsupported or blank answer areas are not part of the {cm['eligible']:,} components.")
    p(f"The candidate also contains {extra} populated components with no eligible comparison, including components in {extra_rows} additional goods row. These overlapping counts are reported separately and earn no correctness credit. Header and goods components are counted by distinct scope and field keys; unscored output must not be treated as validated content.")
    heading("Timing is limited to observed reading intervals")
    p("For the first Sarens collection, seven successful saved page readings span 2 min 56 sec from first request to last response; their summed page intervals are 4 min 31 sec because work overlaps. Some source proposals still fail validation, so this is not a time to finish an approved form. The broader batch changed worker count while retaining a shared six-request-per-minute cap, included repair requests and administrative resumes, and is not a production speed guarantee.")
    heading("What must happen before claiming 90 percent")
    p("Resolve missing source-backed mappings and item alignment, supply approved business or product context for genuinely absent information, and validate customs-specific decisions separately. Then evaluate the unchanged rules on the 11 held-out collections and obtain independent review of content equivalence and extra output. The present development score cannot establish performance on new shipments or complete coverage of the three forms.")
    p("Under the requested workflow, unsupported form values stay blank and a draft email requests the missing details. The email remains unsent. No invented defaults receive credit; these blanks still count against field coverage and comparable-component agreement.")
    heading("Frozen evidence")
    p("The report is generated from an immutable paired snapshot containing both exact reports, the supplementary report, review records, development/holdout design, coverage audit and code hashes. It preserves the earlier benchmark Word report.")
    p("Snapshot folder: " + snapshot.name, small=True)
    p("Supplementary report SHA256: " + hashlib.sha256((snapshot / "content-agreement-report.json").read_bytes()).hexdigest(), small=True)
    p("The manifest records full SHA256 hashes for every frozen source and the original exact scorer and prompt. Original readings were reused; this report does not represent a new extraction run.", small=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    d.save(output)
    print(json.dumps({"output": str(output), "snapshot": str(snapshot), "exactAgreement": cm["exactAgreement"], "contentAgreement": cm["contentAgreement"]}))


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--snapshot", type=pathlib.Path, required=True)
    ap.add_argument("--output", type=pathlib.Path, required=True)
    args = ap.parse_args()
    create(args.snapshot, args.output)
