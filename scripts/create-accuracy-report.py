#!/usr/bin/env python3
"""Create a frozen, reviewable Word report from a benchmark JSON snapshot.

Use the bundled Codex Python runtime. This reads only the supplied snapshot;
it never calls the provider or updates the live benchmark/workspace.
"""
import argparse
import datetime as dt
import hashlib
import json
import pathlib

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def pct(value):
    return "Not available" if value is None else f"{100 * value:.2f}%"


def integer(value):
    return f"{int(value):,}"


def duration(milliseconds):
    if milliseconds is None:
        return "Not available"
    seconds = round(milliseconds / 1000)
    return f"{seconds // 60} min {seconds % 60:02d} sec"


def put(parent, tag, attrs):
    element = OxmlElement(tag)
    for name, value in attrs.items():
        element.set(qn(name), str(value))
    parent.append(element)
    return element


def create(snapshot, output, background_paused=True):
    raw = snapshot.read_bytes()
    report = json.loads(raw)
    checksum = hashlib.sha256(raw).hexdigest()
    summary = report["summary"]
    result = summary["allAttempted"]
    progress = report["progress"]
    terminal = bool(progress.get("terminal"))
    scored = [p for p in report["packets"] if p.get("metrics") is not None]
    frozen = dt.datetime.fromisoformat(report["generatedAt"].replace("Z", "+00:00"))
    moment = frozen.strftime("%d %B %Y at %H:%M:%S UTC")
    filled = result["match"] + result["mismatch"]
    eligible = summary["eligiblePackets"]
    unfinished = eligible - summary["scoredPackets"]
    actual_failures = [(p["folder"], f) for p in report["packets"] for f in p.get("pipelineFailures", []) if f.get("error") != "Not run."]
    failure_kinds = sorted(set(f.get("error", "") for _, f in actual_failures))
    label_conflicts = sum(len(p.get("labelConflicts", [])) for p in report["packets"])
    incomplete_prints = sum(any("Goods score excluded" in w for w in p.get("warnings", [])) for p in report["packets"])

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.66)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.footer_distance = Inches(0.3)
    styles = document.styles
    for border in styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Calibri")
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Title"].font.size = Pt(26)
    styles["Title"].font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(8)
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(9)
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)
    document.core_properties.title = "AICD shipment form accuracy report"
    document.core_properties.subject = "Interim benchmark against employee ECUS declarations"
    document.core_properties.author = "AICD"
    document.core_properties.keywords = "AICD, Kimi, shipment, accuracy, interim benchmark"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    frun = footer.add_run("AICD accuracy report  |  ")
    frun.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    def paragraph(text, lead=None, small=False):
        p = document.add_paragraph()
        if lead:
            p.add_run(lead + " ").bold = True
        p.add_run(text)
        if small:
            for run in p.runs:
                run.font.size = Pt(9)
        return p

    def heading(text, level=1):
        return document.add_heading(text, level)

    def newpage(text):
        p = heading(text)
        p.paragraph_format.page_break_before = True

    def table(headers, rows, widths, font_size=10):
        t = document.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for c, width in zip(t.columns, widths):
            c.width = Inches(width)
        borders = put(t._tbl.tblPr, "w:tblBorders", {})
        for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            put(borders, "w:" + edge, {"w:val": "single", "w:sz": 4, "w:color": "D9D9D9"})
        for i, row in enumerate([headers] + rows):
            cells = t.rows[0].cells if i == 0 else t.add_row().cells
            trPr = t.rows[i]._tr.get_or_add_trPr()
            put(trPr, "w:cantSplit", {})
            if i == 0:
                put(trPr, "w:tblHeader", {})
            for j, (cell, text, width) in enumerate(zip(cells, row, widths)):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr = cell._tc.get_or_add_tcPr()
                margins = put(pr, "w:tcMar", {})
                for edge in ["top", "bottom"]:
                    put(margins, "w:" + edge, {"w:w": 95, "w:type": "dxa"})
                for edge in ["left", "right"]:
                    put(margins, "w:" + edge, {"w:w": 110, "w:type": "dxa"})
                fill = "DCE6EE" if i == 0 else ("F5F7F9" if i % 2 == 0 else "FFFFFF")
                put(pr, "w:shd", {"w:fill": fill})
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.02
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(str(text))
                run.bold = i == 0
                run.font.size = Pt(font_size)
                run.font.color.rgb = RGBColor(0, 0, 0)
        document.add_paragraph().paragraph_format.space_after = Pt(0)
        return t

    document.add_paragraph("AICD shipment form accuracy report", "Title")
    paragraph(("Benchmark measurement" if terminal else "Interim measurement") + " against employee ECUS declarations", small=False)
    paragraph("Results frozen on " + moment, small=True)
    status_text = "The page-reading run has ended, but the drafts still require review." if terminal else "The batch is still running, and the current drafts require review. This is not a final full-batch accuracy result."
    paragraph(f"Exact agreement of comparable draft components is {pct(result['accuracy'])}, measured against employee answers for {len(scored)} of {eligible} eligible import collections. {status_text} This comparison does not establish independently verified correctness or a time to finish an approved shipment form.")
    heading("Measured results", 2)
    table(["Metric", "Calculation", "Current result"], [
        ["Exact answer agreement", f"{integer(result['match'])} / {integer(result['total'])}", pct(result["accuracy"])],
        ["Filled value agreement", f"{integer(result['match'])} / {integer(filled)}", pct(result["filledPrecision"])],
        ["Comparable field coverage", f"{integer(filled)} / {integer(result['total'])}", pct(result["coverage"])],
    ], [3.05, 1.8, 2.15], 11)
    paragraph(f"The scored answer set contains {integer(result['match'])} matches, {integer(result['mismatch'])} mismatches and {integer(result['missing'])} missing values. Missing values count against exact agreement, including values on pages that have not yet been read. Filled value agreement covers only fields with an employee answer to compare.")
    heading("How much of the test is covered", 2)
    table(["Population", "At this snapshot"], [
        ["All requested sample collections", "44 collections"],
        ["Prepared original sources", "223 files and 497 rendered pages"],
        ["Eligible import collections with supported answers", f"{eligible} collections and {integer(summary['eligibleComponents'])} answer components"],
        ["Eligible collections currently scored", f"{len(scored)} / {eligible} ({pct(len(scored)/eligible)})"],
        ["Eligible answer components represented in scores", f"{integer(result['total'])} / {integer(summary['eligibleComponents'])} ({pct(summary['scoredSampleComponentCoverage'])})"],
        ["Successful saved page readings across the batch", f"{progress['successfulPages']} / {progress['totalPages']} ({pct(progress['successfulPages']/progress['totalPages'])})"],
    ], [4.1, 2.9], 10)
    paragraph(f"{summary['provisionalPackets']} scored collections remain provisional; {summary['completedReadingPackets']} have complete validated readings. The completed-reading-only exact agreement is {pct(summary['completedOnly']['accuracy']).lower()}. The {unfinished} other eligible collections have no comparable score" + ("." if terminal else " yet."), lead="Incomplete coverage.")

    newpage("What the metrics measure")
    heading("Definitions and formulas", 2)
    paragraph("Matches divided by all eligible answer components. The denominator includes matches, mismatches and missing values in collections with at least one successful provider response. This is the report's primary accuracy measure.", lead="Exact answer agreement.")
    paragraph("Matches divided by matches plus mismatches. This asks how often populated, comparable values agree with the employee. It excludes blanks and output without a comparable answer, so it cannot establish that every populated field is correct.", lead="Filled value agreement.")
    paragraph("Matches plus mismatches, divided by all eligible answer components. This measures how much of the comparable answer set the draft has populated.", lead="Comparable field coverage.")
    paragraph("Eligible components represented by currently scored collections divided by all 2,316 eligible components. This measures test coverage, not the percentage of fields the AI has filled.", lead="Scored sample coverage.")
    paragraph("The score counts individual components, such as quantity and unit, rather than complete fields or complete forms. Longer declarations therefore carry more weight. Goods are compared in printed row order; an extra or missing row can shift later comparisons.")
    heading("Benchmark method", 2)
    paragraph("The pipeline receives original commercial source files. Visual documents are rendered to page images for Kimi. Employee ToKhai files and other declaration outputs are held out and never supplied to Kimi; they are read locally only for scoring. Saved historical AI readings and draft values are excluded from the benchmark inputs.")
    paragraph("The audited sample folder supplies the document partition and import or export direction. This is a test fixture assumption. Automatic grouping, email delivery, and assignment of documents to real shipments are outside this accuracy test.")
    paragraph("Comparison uses mapped, nonblank employee answer components that agree across available revisions. It normalizes case, whitespace and explicit numeric formatting, then requires exact agreement. It does not judge semantic equivalence, translate English descriptions into Vietnamese, or verify customs classification. Conflicting employee revisions and unsupported answer formats are excluded.")
    heading("Limits of the answer key", 2)
    paragraph("An employee mismatch is a difference to investigate, not independent proof that the AI is wrong. The employee may have used information absent from these sources, a translated description, or a customs decision. Extra draft fields and extra goods rows are reported separately; they are not validated by the score.")
    paragraph("A duplicate audit of the 29 eligible collections found 40 distinct supported employee-file hashes and 152 distinct source-file hashes, with no identical source sets or answer files shared across eligible folders. Distinct files still share customers and templates, so they do not establish statistical independence.")

    newpage("Current results and exceptions")
    heading("Results by answer type", 2)
    table(["Answer type", "Matches", "Mismatch", "Missing", "Exact agreement"], [
        [label, integer(result[key]["match"]), integer(result[key]["mismatch"]), integer(result[key]["missing"]), pct(result[key]["accuracy"])]
        for label, key in [("Shipment headers", "header"), ("Goods components", "goods")]
    ], [2.4, 1.0, 1.1, 1.1, 1.4], 10)
    heading("Collection detail", 2)
    displayed = scored[:8]
    if len(scored) > len(displayed):
        paragraph(f"The table shows the first {len(displayed)} of {len(scored)} scored collections in inventory order. All {len(scored)} contribute to the aggregate results.", small=True)
    table(["Collection", "Matches / answers", "Exact", "Filled coverage"], [
        [p["folder"], f"{p['metrics']['match']} / {p['metrics']['total']}", pct(p['metrics']['accuracy']), pct(p['metrics']['coverage'])] for p in displayed
    ], [3.55, 1.25, 1.05, 1.15], 9.5)
    paragraph(f"The {summary['provisionalPackets']} provisional scores use the full eligible denominator for their collections. A successful page response can still contain rejected or ambiguous proposals. " + ("Unresolved readings remain incomplete after the run." if terminal else "Scores can change as remaining pages are read."), small=True)
    heading("Failures and exclusions", 2)
    paragraph(f"Across the full batch, {progress['failedPages']} page outcome(s) failed and {progress['pendingPages']} pages remain pending. Preparation failures: {progress['preparationFailures']}. A saved reading may involve several provider requests, including repair. Rejected proposals are separate review issues.")
    if any("timed out" in message.lower() for message in failure_kinds):
        paragraph("A page timed out in 2015/SARENS-02pkg-91kg-SCSC. Its proposals were not saved, and it was not automatically retried.")
    paragraph("Fifteen collections are outside the import score: 4 lack primary sources, 9 are non-import collections, and 2 have no supported import answer key. All remain visible in the sample inventory.")
    paragraph(f"The score excludes {label_conflicts} conflicting answer-component keys. Goods scores for {incomplete_prints} incomplete employee prints are excluded pending review of missing printed blocks; their eligible headers remain comparable.")
    paragraph(f"The scored drafts have {result['extraGoodsRows']} extra goods row(s) in {result['packetsWithExtraGoodsRows']} collection(s), plus {result['unscoredPopulatedComponents']} distinct populated components without comparable answers. These outputs are outside the score and unvalidated, not a count of errors.")

    newpage("What the first Sarens collection reveals")
    first = next(p for p in report['packets'] if p['folder'] == '01x40FR-Sarens')
    fm = first['metrics']
    paragraph(f"We visually reviewed 01x40FR-Sarens. Its exact agreement is {fm['match']} / {fm['total']} ({pct(fm['accuracy'])}); filled value agreement is {pct(fm['filledPrecision'])}, with {pct(fm['coverage'])} coverage. All seven pages had successful responses. Four rejected proposals across three pages and unresolved review issues keep it incomplete. These are concrete extraction or reconciliation findings.")
    heading("Accessory package becomes an extra goods row", 2)
    paragraph("The invoice lists two power packs at EUR 12,500 each and a third accessory box whose value is marked included. The agreement lists two sets with accessories; the employee declaration likewise includes accompanying tools in its two goods descriptions. The draft retained the accessory package as a third, unpriced goods row. The item is present in the source, but the packaging and declaration rows were not reconciled.")
    heading("Correctly read unit prices do not reach the draft", 2)
    paragraph("Agreement page 1 explicitly shows a unit price of EUR 12,500 for each power pack. Kimi extracted both values correctly, but the selected invoice-based draft rows leave both unit prices blank. This is an assembly gap after extraction.")
    heading("Date and address roles are confused", 2)
    paragraph("The sea waybill's shipped-on-board date, 12 December 2016, was proposed as an arrival date. Validation withheld it from the form, but its field role is wrong. Separately, the agreement's exporter telephone and fax were accepted as a second address line and appeared in the draft address.")
    heading("Payment scope and harmless formatting need review", 2)
    paragraph("The agreement allows free use of the equipment and separately describes transfer payments for maintenance and operator allowances. The draft payment-method field contains that allowances paragraph. The system also treats CO., LTD versus CO. LTD and spaced versus unspaced container identifiers as conflicts. A contract allowing any port in Belgium conflicts with the actual waybill port of Antwerp.")
    heading("Some employee differences need additional evidence", 2)
    paragraph("The employee HS code is 84089050, while the invoice prints 8709 9000. The employee records Cat Lai as the unloading port, while the waybill prints Vung Tau. These differences cannot be labelled reading errors solely from the answer comparison. They require business reconciliation or further documents.")
    paragraph("Evidence reviewed: Invoice & Packing list - Power pack (T.import).xlsx, pages 1 and 2; Borrow agreement - Power pack.docx, pages 1 and 2; Seaway BL - MCOP0101_332294764.pdf, its single supplied page; and held-out ToKhaiHQ7N_QDTQ_101238157201.xls. The first collection is an illustrative audit, not a representative error-rate sample.", small=True)

    newpage("Processing time and automatic operation")
    heading("Observed time is not form completion time", 2)
    paragraph(f"The median first-request-to-last-response span among {len(scored)} scored collections is {duration(result['medianObservedResponseSpanMs'])}. It includes waiting, interruptions and resumed cached work. For partial collections it ends at the latest recorded response, not at a complete or approved shipment form.")
    paragraph(f"The seven pages in 01x40FR-Sarens took {duration(first['timing']['firstStartToLastFinishMs'])} from first request to last response. The sum of recorded page intervals was {duration(first['timing']['providerWorkMs'])}. These intervals can include validation repair and rate-limit waits; they are not pure compute or billed work. Overlap makes their sum longer than elapsed time. The draft still needed review afterwards.")
    paragraph(f"Concurrency changed from 2 to 4 to {progress.get('benchmarkConcurrency', 8)} workers. There were {progress.get('administrativelyInterruptedPages', 0)} recorded interrupted pages and an earlier restart with at most {progress.get('earlierAbruptRestartInFlightAtMost', 0)} in-flight requests. This is not a first-attempt-only run or a fixed-concurrency service-level test.")
    heading("Automatic processing now available", 2)
    paragraph("The local application has a durable background queue for imported email attachments and uploads, with parallel processing and safe saving against the latest workspace state. Jobs continue without selecting a shipment while the local server remains running. Failed reads and uncertain results remain visible for retry or review; automatic processing does not approve declarations.")
    paragraph("Documents shows original source pages and file downloads. Settings shows the current Kimi prompts, queue progress, failures and the duration of the most recently completed page.")
    paragraph("Local background processing is paused during the benchmark. Automatic resumption is configured after the full run, for the approved pending work. The Mac must remain awake and running; the handoff does not survive a reboot." if background_paused else "Local background processing is enabled separately from the benchmark run. The local server and Mac must remain running.")
    paragraph("This change processes imported email files and uploads after they enter the application. Gmail and live inbox delivery are not connected. Folder grouping, user review, ECUS submission and approved form completion are not measured by this benchmark.")
    heading("Next measurement", 2)
    paragraph(("Review the remaining failures and rejected proposals, then update the comparison." if terminal else "Complete the remaining pages, review failures and rejected proposals, then refresh the full batch comparison.") + " Separate semantic reading errors from source inconsistencies and customs decisions. Measure one complete shipment from ingestion through review under fixed concurrency before setting an expected completion time.")
    if not terminal:
        paragraph("An automatic follow-up is scheduled to refresh the report after the batch reaches a terminal state.", small=True)
    heading("Result provenance", 2)
    paragraph(f"Frozen result time: {moment}. Model: {progress['model']}. Source snapshot: {snapshot.name}. " + ("The snapshot records the terminal run status." if terminal else "The live benchmark continues independently, so later responses are outside this document."), small=True)
    paragraph("Snapshot SHA256: " + checksum, small=True)
    paragraph("Supporting records: benchmark manifest, per-collection results and original rendered pages. Comparison logic: scripts/score-tokhai-benchmark.py.", small=True)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(json.dumps({"output": str(output), "snapshot": str(snapshot), "snapshotGeneratedAt": report['generatedAt'], "snapshotSha256": checksum, "metrics": result, "progress": {k: progress[k] for k in ['successfulPages', 'failedPages', 'pendingPages', 'totalPages']}}, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--snapshot", required=True, type=pathlib.Path)
    parser.add_argument("--output", required=True, type=pathlib.Path)
    parser.add_argument("--background-enabled", action="store_true")
    args = parser.parse_args()
    create(args.snapshot.resolve(), args.output.resolve(), not args.background_enabled)
